import io
from datetime import datetime
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from sqlalchemy import func
from database.models import Transaction, Category, TransactionType

async def get_financial_summary_by_period(user_id: int, start_date: datetime, end_date: datetime, session: AsyncSession):
    """Davr bo'yicha jami kirim, chiqim va qoldiqni hisoblash"""
    income_stmt = select(func.coalesce(func.sum(Transaction.amount), 0.0)).where(
        Transaction.user_id == user_id,
        Transaction.type == TransactionType.INCOME,
        Transaction.created_at >= start_date,
        Transaction.created_at <= end_date
    )
    total_income = (await session.execute(income_stmt)).scalar()

    expense_stmt = select(func.coalesce(func.sum(Transaction.amount), 0.0)).where(
        Transaction.user_id == user_id,
        Transaction.type == TransactionType.EXPENSE,
        Transaction.created_at >= start_date,
        Transaction.created_at <= end_date
    )
    total_expense = (await session.execute(expense_stmt)).scalar()

    return {
        "income": total_income,
        "expense": total_expense,
        "balance": total_income - total_expense
    }

async def get_expenses_by_period(user_id: int, start_date: datetime, end_date: datetime, session: AsyncSession):
    """Kategoriyalar bo'yicha xarajatlar yig'indisini olish (Diagramma uchun)"""
    stmt = (
        select(Category.name, func.sum(Transaction.amount).label("total"))
        .join(Category, Transaction.category_id == Category.id)
        .where(
            Transaction.user_id == user_id,
            Transaction.type == TransactionType.EXPENSE,
            Transaction.created_at >= start_date,
            Transaction.created_at <= end_date
        )
        .group_by(Category.name)
    )
    result = await session.execute(stmt)
    return result.all()

async def get_detailed_transactions(user_id: int, start_date: datetime, end_date: datetime, session: AsyncSession):
    """Excel uchun barcha tranzaksiyalarni olish"""
    stmt = (
        select(
            Transaction.created_at,
            Transaction.type,
            Category.name.label("category"),
            Transaction.amount,
            Transaction.description
        )
        .outerjoin(Category, Transaction.category_id == Category.id)
        .where(
            Transaction.user_id == user_id,
            Transaction.created_at >= start_date,
            Transaction.created_at <= end_date
        )
        .order_by(Transaction.created_at.desc())
    )
    result = await session.execute(stmt)
    return result.all()

def generate_pie_chart(data: list[tuple[str, float]]) -> io.BytesIO:
    """Xarajatlar strukturasi uchun doiraviy grafik (Pie Chart)"""
    labels = [row[0] for row in data]
    amounts = [row[1] for row in data]

    fig, ax = plt.subplots(figsize=(6, 6))
    ax.pie(amounts, labels=labels, autopct='%1.1f%%', startangle=140)
    ax.set_title("Xarajatlar strukturasi")

    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight')
    buf.seek(0)
    plt.close(fig)
    return buf

def generate_excel_report(transactions: list, summary: dict) -> io.BytesIO:
    """1-rasmdagi ko'rinishga 100% moslashtirilgan chiroyli Excel jadvali"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws.views.sheetView[0].showGridLines = True

    # Standart chegaralar (chiziqlar)
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    # 1. Sarlavha qatori
    headers = ["Sana va vaqt", "Turi", "Kategoriya", "Summa (so'm)", "Izoh"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # 2. Ma'lumot qatorlari
    current_row = 2
    for row in transactions:
        created_at, trans_type, category, amount, description = row
        type_str = "Kirim" if trans_type == TransactionType.INCOME else "Chiqim"
        date_str = created_at.strftime("%Y-%m-%d %H:%M:%S")
        cat_str = category if category else "Noma'lum"
        desc_str = description if description else "-"

        ws.append([date_str, type_str, cat_str, amount, desc_str])

        # Har bir katakni formatlash
        ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=3).alignment = Alignment(horizontal="left", vertical="center")
        
        sum_cell = ws.cell(row=current_row, column=4)
        sum_cell.alignment = Alignment(horizontal="right", vertical="center")
        sum_cell.number_format = '#,##0'

        ws.cell(row=current_row, column=5).alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, 6):
            ws.cell(row=current_row, column=col_idx).border = thin_border

        current_row += 1

    # 3. Yakuniy qatorlar (Jami Kirim, Jami Chiqim, Sof Qoldiq)
    summary_rows = [
        ("JAMI KIRIM", "Kirim", "", summary["income"], ""),
        ("JAMI CHIQIM", "Chiqim", "", summary["expense"], ""),
        ("SOF QOLDIQ", "Qoldiq", "", summary["balance"], "")
    ]

    summary_font = Font(name="Calibri", size=11, bold=True)

    for title, t_type, cat, amount, note in summary_rows:
        ws.append([title, t_type, cat, amount, note])

        ws.cell(row=current_row, column=1).font = summary_font
        ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

        ws.cell(row=current_row, column=2).font = summary_font
        ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=current_row, column=3).alignment = Alignment(horizontal="left", vertical="center")

        sum_cell = ws.cell(row=current_row, column=4)
        sum_cell.font = summary_font
        sum_cell.alignment = Alignment(horizontal="right", vertical="center")
        sum_cell.number_format = '#,##0'

        ws.cell(row=current_row, column=5).alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, 6):
            ws.cell(row=current_row, column=col_idx).border = thin_border

        current_row += 1

    # 4. Ustun kengliklarini avtomatik hisoblash (Auto-fit Columns)
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        # Matn sig'ishi uchun qo'shimcha bo'sh joy qo'shamiz
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf